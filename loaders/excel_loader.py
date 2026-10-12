from langchain_core.documents import Document
import openpyxl
import os


def load_excel(folder):
    documents = []

    for file in os.listdir(folder):

        if file.endswith(".xlsx"):

            path = os.path.join(folder, file)

            workbook = openpyxl.load_workbook(path, data_only=True)

            for sheet in workbook.worksheets:

                rows = list(sheet.iter_rows(values_only=True))

                if len(rows) < 2:
                    continue

                headers = rows[0]

                for row in rows[1:]:

                    text = ""

                    for h, v in zip(headers, row):

                        text += f"{h}：{v}\n"

                    documents.append(
                        Document(
                            page_content=text,
                            metadata={
                                "source": file,
                                "sheet": sheet.title
                            }
                        )
                    )

    print(f"读取到 {len(documents)} 条Excel记录")

    return documents
